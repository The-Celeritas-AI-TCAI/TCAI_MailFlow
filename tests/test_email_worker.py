import os
import importlib.util
import smtplib
import sys
import tempfile
import threading
import time
import types
import unittest
from unittest.mock import patch

# Keep the worker unit test runnable without requiring the web application's
# optional environment-loader package to be installed in the test interpreter.
dotenv_stub = types.ModuleType("dotenv")
dotenv_stub.load_dotenv = lambda: None
sys.modules.setdefault("dotenv", dotenv_stub)

import database
from reporting import build_report
from config import Config
from config import normalized_smtp_username
from email_worker import CampaignManager, EmailMessageBuilder, ResilientSMTP, SMTPConnectionManager


class CampaignFailureIsolationTest(unittest.TestCase):
    def setUp(self):
        self.temp_db = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
        self.temp_db.close()
        self.original_db = Config.DATABASE_URL
        self.original_retries = Config.MAX_EMAIL_RETRIES
        self.original_username = Config.SMTP_USERNAME
        self.original_password = Config.SMTP_PASSWORD
        self.original_health_check_idle = Config.SMTP_HEALTH_CHECK_IDLE
        self.original_interval_seconds = Config.EMAIL_INTERVAL_SECONDS
        Config.DATABASE_URL = f"sqlite:///{self.temp_db.name}"
        Config.MAX_EMAIL_RETRIES = 1
        Config.EMAIL_INTERVAL_SECONDS = 0
        Config.SMTP_USERNAME = "sender@example.test"
        Config.SMTP_PASSWORD = "test-password"
        Config.SMTP_HEALTH_CHECK_IDLE = 30
        database.init_database()

    def tearDown(self):
        Config.DATABASE_URL = self.original_db
        Config.MAX_EMAIL_RETRIES = self.original_retries
        Config.SMTP_USERNAME = self.original_username
        Config.SMTP_PASSWORD = self.original_password
        Config.SMTP_HEALTH_CHECK_IDLE = self.original_health_check_idle
        Config.EMAIL_INTERVAL_SECONDS = self.original_interval_seconds
        # The daemon worker can briefly retain SQLite's Windows file handle after
        # queue.join(); a failed cleanup must not invalidate the behavior test.
        try:
            os.unlink(self.temp_db.name)
        except PermissionError:
            pass

    def test_disconnect_for_one_recipient_does_not_block_the_next(self):
        campaign_id = database.create_campaign(
            ["first@example.test", "broken@example.test", "third@example.test"],
            "Subject", "Body", [],
        )
        recipient_ids = database.queued_recipient_ids(campaign_id)
        manager = CampaignManager()
        calls = []

        def fake_send(_, __, recipient, *args):
            calls.append(recipient)
            if recipient == "broken@example.test":
                raise smtplib.SMTPServerDisconnected("forced disconnect")

        with patch("email_worker.SMTPConnectionManager.send_message", new=fake_send):
            worker = threading.Thread(target=manager._worker, daemon=True)
            worker.start()
            for recipient_id in recipient_ids:
                manager.recipient_queue.put(recipient_id)
            manager.recipient_queue.join()

        campaign = database.get_campaign(campaign_id)
        self.assertEqual(campaign["sent"], 2)
        self.assertEqual(campaign["retrying"], 1)
        self.assertEqual(campaign["status"], "running")
        self.assertEqual(calls.count("broken@example.test"), 1)
        self.assertIn("third@example.test", calls)

    def test_new_attachment_rows_store_only_filesystem_metadata(self):
        attachment = tempfile.NamedTemporaryFile(delete=False)
        attachment.write(b"campaign attachment")
        attachment.close()
        try:
            campaign_id = database.create_campaign(
                ["recipient@example.test"], "Subject", "Body",
                [{"filename": "attachment.txt", "file_path": attachment.name, "size": 19, "mime_type": "text/plain"}],
            )
            with database.connection() as db:
                columns = [row[1] for row in db.execute("PRAGMA table_info(attachments)")]
                row = db.execute("SELECT filename,file_path,size,mime_type FROM attachments WHERE campaign_id=?", (campaign_id,)).fetchone()
            self.assertNotIn("content", columns)
            self.assertEqual(dict(row)["file_path"], attachment.name)
        finally:
            os.unlink(attachment.name)

    def test_timeout_for_one_recipient_does_not_block_the_next(self):
        campaign_id = database.create_campaign(
            ["first@example.test", "timeout@example.test", "third@example.test"],
            "Subject", "Body", [],
        )
        manager, calls = CampaignManager(), []

        def fake_send(_, __, recipient, *args):
            calls.append(recipient)
            if recipient == "timeout@example.test":
                raise TimeoutError("forced timeout")

        with patch("email_worker.SMTPConnectionManager.send_message", new=fake_send):
            threading.Thread(target=manager._worker, daemon=True).start()
            for recipient_id in database.queued_recipient_ids(campaign_id):
                manager.recipient_queue.put(recipient_id)
            manager.recipient_queue.join()

        campaign = database.get_campaign(campaign_id)
        self.assertEqual((campaign["sent"], campaign["retrying"], campaign["status"]), (2, 1, "running"))
        self.assertEqual(calls.count("timeout@example.test"), 1)
        self.assertIn("third@example.test", calls)

    def test_transient_failures_stop_after_the_persisted_retry_limit(self):
        campaign_id = database.create_campaign(["offline@example.test"], "Subject", "Body", [])
        manager, calls = CampaignManager(), []

        def fake_send(_, __, recipient, *args):
            calls.append(recipient)
            raise smtplib.SMTPServerDisconnected("forced disconnect")

        with patch("email_worker.SMTPConnectionManager.send_message", new=fake_send):
            threading.Thread(target=manager._worker, daemon=True).start()
            manager.recipient_queue.put(database.queued_recipient_ids(campaign_id)[0])
            manager.recipient_queue.join()
            with database.connection() as db:
                db.execute("UPDATE recipients SET status='queued',next_attempt_at=NULL WHERE campaign_id=?", (campaign_id,))
            manager.recipient_queue.put(database.queued_recipient_ids(campaign_id)[0])
            manager.recipient_queue.join()

        campaign = database.get_campaign(campaign_id)
        self.assertEqual(calls, ["offline@example.test", "offline@example.test"])
        self.assertEqual((campaign["failed"], campaign["retrying"], campaign["status"]), (1, 0, "failed"))
        with database.connection() as db:
            self.assertEqual(db.execute("SELECT attempts FROM recipients WHERE campaign_id=?", (campaign_id,)).fetchone()[0], 2)

    def test_attachment_campaign_builds_and_sends_each_recipient(self):
        attachment = tempfile.NamedTemporaryFile(suffix=".txt", delete=False)
        attachment.write(b"attachment content")
        attachment.close()
        try:
            campaign_id = database.create_campaign(
                ["one@example.test", "two@example.test", "three@example.test"], "Subject", "Body",
                [{"filename": "attachment.txt", "file_path": attachment.name, "size": 18, "mime_type": "text/plain"}],
            )
            manager, calls = CampaignManager(), []

            def fake_send(_, message_bytes, recipient, has_attachments=False):
                self.assertTrue(has_attachments)
                self.assertIn(b"attachment.txt", message_bytes)
                calls.append(recipient)

            with patch("email_worker.SMTPConnectionManager.send_message", new=fake_send):
                threading.Thread(target=manager._worker, daemon=True).start()
                for recipient_id in database.queued_recipient_ids(campaign_id):
                    manager.recipient_queue.put(recipient_id)
                manager.recipient_queue.join()

            campaign = database.get_campaign(campaign_id)
            self.assertEqual((campaign["sent"], campaign["failed"], campaign["status"]), (3, 0, "completed"))
            self.assertEqual(calls, ["one@example.test", "two@example.test", "three@example.test"])
        finally:
            os.unlink(attachment.name)

    def test_sender_rejection_fails_campaign_without_retrying_each_recipient(self):
        campaign_id = database.create_campaign(
            ["one@example.test", "two@example.test", "three@example.test"], "Subject", "Body", [],
        )
        manager, calls = CampaignManager(), []

        def fake_send(_, __, recipient, *args):
            calls.append(recipient)
            raise smtplib.SMTPSenderRefused(501, b"syntax error", "sender@example.test")

        with patch("email_worker.SMTPConnectionManager.send_message", new=fake_send):
            threading.Thread(target=manager._worker, daemon=True).start()
            for recipient_id in database.queued_recipient_ids(campaign_id):
                manager.recipient_queue.put(recipient_id)
            manager.recipient_queue.join()

        campaign = database.get_campaign(campaign_id)
        self.assertEqual((campaign["sent"], campaign["failed"], campaign["status"]), (0, 3, "failed"))
        self.assertEqual(calls, ["one@example.test"])
        self.assertIn("SMTP sender address rejected", campaign["fatal_error"])

    def test_send_message_uses_bare_envelope_sender_and_correct_stages(self):
        class FakeSMTP:
            def __init__(self): self.sock = None; self.calls = []
            def has_extn(self, name): return False
            def mail(self, sender, options=()): self.calls.append(("MAIL_FROM", sender, options)); return 250, b"ok"
            def rcpt(self, recipient): self.calls.append(("RCPT_TO", recipient)); return 250, b"ok"
            def data(self, message): self.calls.append(("DATA", message)); return 250, b"ok"

        connection = SMTPConnectionManager()
        connection.smtp = FakeSMTP()
        connection.send_message(b"message", "recipient@example.test")
        self.assertEqual(connection.smtp.calls[0], ("MAIL_FROM", "sender@example.test", []))
        self.assertEqual([call[0] for call in connection.smtp.calls], ["MAIL_FROM", "RCPT_TO", "DATA"])

    def test_idle_connection_is_checked_then_reused(self):
        class FakeSMTP:
            sock = None
            def __init__(self): self.calls = []
            def noop(self): self.calls.append("NOOP"); return 250, b"ok"
            def has_extn(self, name): return False
            def mail(self, sender, options=()): self.calls.append("MAIL_FROM"); return 250, b"ok"
            def rcpt(self, recipient): self.calls.append("RCPT_TO"); return 250, b"ok"
            def data(self, message): self.calls.append("DATA"); return 250, b"ok"

        connection = SMTPConnectionManager()
        connection.smtp = FakeSMTP()
        connection.last_activity = time.monotonic() - 31
        connection.send_message(b"message", "recipient@example.test")
        self.assertEqual(connection.smtp.calls, ["NOOP", "MAIL_FROM", "RCPT_TO", "DATA"])

    def test_attachment_mime_part_appears_once_and_round_trips_original_bytes(self):
        attachment_bytes = os.urandom(100 * 1024)
        builder = EmailMessageBuilder()
        message = builder.build(
            "sender@example.test", "recipient@example.test", "Subject", "Body",
            [{"filename": "sample.bin", "content": attachment_bytes}],
        )
        parts = list(message.iter_attachments())
        self.assertEqual(len(parts), 1)
        self.assertEqual(parts[0].get_filename(), "sample.bin")
        self.assertEqual(parts[0].get_payload(decode=True), attachment_bytes)
        self.assertEqual(parts[0]["Content-Transfer-Encoding"], "base64")
        self.assertIn(b"MIME-Version: 1.0", message.as_bytes())

    def test_attachment_size_threshold_mime_builds(self):
        builder = EmailMessageBuilder()
        for raw_size in (50 * 1024, 500 * 1024, 1024 * 1024):
            raw = b"x" * raw_size
            message = builder.build(
                "sender@example.test", "recipient@example.test", "Subject", "Body",
                [{"filename": "threshold.bin", "content": raw}],
            )
            serialized = message.as_bytes()
            part = next(message.iter_attachments())
            self.assertEqual(part.get_payload(decode=True), raw)
            self.assertGreater(len(serialized), raw_size)

    def test_resilient_smtp_chunks_large_data_writes(self):
        class RecordingSocket:
            def __init__(self): self.writes = []
            def gettimeout(self): return 30
            def send(self, data): self.writes.append(bytes(data)); return len(data)

        smtp = ResilientSMTP.__new__(ResilientSMTP)
        smtp.sock = RecordingSocket()
        smtp.send(b"x" * (ResilientSMTP.TLS_WRITE_CHUNK_SIZE * 2 + 1))
        self.assertEqual([len(write) for write in smtp.sock.writes], [65536, 65536, 1])

    def test_sender_rejection_is_identified_as_mail_from(self):
        class RejectingSMTP:
            sock = None
            def has_extn(self, name): return False
            def mail(self, sender, options=()): return 501, b"syntax error in parameters or arguments"

        connection = SMTPConnectionManager()
        connection.smtp = RejectingSMTP()
        with self.assertRaises(smtplib.SMTPSenderRefused):
            connection.send_message(b"message", "recipient@example.test")
        self.assertEqual(connection.stage, "MAIL_FROM")

    def test_smtp_username_rejects_display_name_and_control_characters(self):
        self.assertEqual(normalized_smtp_username(" sender@example.test "), "sender@example.test")
        for invalid in ('Name <sender@example.test>', 'sender@example.test,other@example.test', 'sender@\nexample.test'):
            with self.assertRaises(ValueError):
                normalized_smtp_username(invalid)

    def test_recovery_pauses_interrupted_campaigns_without_touching_scheduled_campaigns(self):
        interrupted_id = database.create_campaign(["sent@example.test", "interrupted@example.test"], "Subject", "Body", [])
        future_id = database.create_campaign(
            ["future@example.test"], "Subject", "Body", [],
            scheduled_at="2999-01-01T00:00:00+00:00", timezone_name="UTC",
        )
        with database.connection() as db:
            db.execute("UPDATE campaigns SET status='running' WHERE id=?", (interrupted_id,))
            recipient_ids = [row[0] for row in db.execute("SELECT id FROM recipients WHERE campaign_id=? ORDER BY id", (interrupted_id,))]
            db.execute("UPDATE recipients SET status='sent' WHERE id=?", (recipient_ids[0],))
            db.execute("UPDATE recipients SET status='sending' WHERE id=?", (recipient_ids[1],))

        recovered = database.recover_campaigns()

        self.assertEqual(recovered, [interrupted_id])
        self.assertEqual(database.get_campaign(interrupted_id)["status"], "paused")
        self.assertEqual(database.queued_recipient_ids(interrupted_id), [recipient_ids[1]])
        self.assertIsNone(database.claim_recipient(recipient_ids[1]))
        with database.connection() as db:
            self.assertEqual(db.execute("SELECT status FROM recipients WHERE id=?", (recipient_ids[0],)).fetchone()[0], "sent")
        self.assertEqual(database.get_campaign(future_id)["status"], "scheduled")

    def test_startup_does_not_enqueue_recovered_campaigns(self):
        manager = CampaignManager()
        with patch("email_worker.database.recover_campaigns", return_value=["old-campaign"]), \
             patch("email_worker.threading.Thread"), \
             patch.object(manager, "enqueue") as enqueue:
            manager.start()
        enqueue.assert_not_called()

    def test_due_scheduled_campaign_is_still_activated_for_the_scheduler(self):
        campaign_id = database.create_campaign(
            ["scheduled@example.test"], "Subject", "Body", [],
            scheduled_at="2000-01-01T00:00:00+00:00", timezone_name="UTC",
        )

        self.assertEqual(database.activate_due_campaigns(), [campaign_id])
        self.assertEqual(database.get_campaign(campaign_id)["status"], "queued")

    def test_resume_and_iteration_history_are_database_backed(self):
        campaign_id = database.create_campaign(
            ["one@example.test", "two@example.test", "three@example.test"],
            "Subject", "Body", [], total_iterations=2,
        )
        ids = database.queued_recipient_ids(campaign_id)
        for recipient_id in ids[:2]:
            self.assertIsNotNone(database.claim_recipient(recipient_id, campaign_id))
            database.record_attempt(recipient_id)
            database.recipient_result(recipient_id, "sent")
        database.pause_campaign(campaign_id)
        self.assertEqual(database.get_campaign(campaign_id)["sent"], 2)
        database.resume_campaign(campaign_id)
        self.assertEqual(database.queued_recipient_ids(campaign_id), [ids[2]])
        self.assertIsNotNone(database.claim_recipient(ids[2], campaign_id))
        database.record_attempt(ids[2])
        database.recipient_result(ids[2], "sent")
        campaign = database.get_campaign(campaign_id)
        self.assertEqual((campaign["current_iteration"], campaign["status"], campaign["queued"]), (2, "queued", 3))
        first_iteration = database.iteration_results(campaign_id, 1)[1]
        self.assertEqual([row["status"] for row in first_iteration], ["sent", "sent", "sent"])

    def test_campaign_interval_gates_claims_and_resume_sets_a_new_slot(self):
        campaign_id = database.create_campaign(
            ["one@example.test", "two@example.test"], "Subject", "Body", [],
        )
        first_id, second_id = database.queued_recipient_ids(campaign_id)
        self.assertEqual(database.dispatchable_recipient_ids(campaign_id), [first_id])
        claimed = database.claim_recipient(first_id, campaign_id)
        self.assertIsNotNone(claimed)
        send_slot = database.begin_recipient_send(first_id, 30)
        self.assertIsNotNone(send_slot)
        self.assertEqual(database.dispatchable_recipient_ids(campaign_id), [])
        self.assertIsNone(database.claim_recipient(second_id, campaign_id))
        database.record_attempt(first_id)
        database.recipient_result(first_id, "sent")

        with database.connection() as db:
            db.execute(
                "UPDATE campaigns SET next_run_at=? WHERE id=?",
                ("2999-01-01T00:00:00+00:00", campaign_id),
            )
        Config.EMAIL_INTERVAL_SECONDS = 30
        database.pause_campaign(campaign_id)
        database.resume_campaign(campaign_id)
        self.assertIsNone(database.claim_recipient(second_id, campaign_id))

        with database.connection() as db:
            db.execute(
                "UPDATE campaigns SET next_run_at=NULL WHERE id=?", (campaign_id,)
            )
        self.assertIsNotNone(database.claim_recipient(second_id, campaign_id))

    def test_pause_before_smtp_authorization_releases_claim(self):
        campaign_id = database.create_campaign(["one@example.test"], "Subject", "Body", [])
        recipient_id = database.queued_recipient_ids(campaign_id)[0]
        self.assertIsNotNone(database.claim_recipient(recipient_id, campaign_id))
        database.pause_campaign(campaign_id)
        self.assertIsNone(database.begin_recipient_send(recipient_id, 30))
        database.release_recipient_claim(recipient_id)
        self.assertEqual(database.queued_recipient_ids(campaign_id), [recipient_id])

    def test_iteration_ledger_prevents_second_success(self):
        campaign_id = database.create_campaign(["one@example.test"], "Subject", "Body", [])
        recipient_id = database.queued_recipient_ids(campaign_id)[0]
        self.assertIsNotNone(database.claim_recipient(recipient_id, campaign_id))
        database.record_attempt(recipient_id)
        database.recipient_result(recipient_id, "sent")
        self.assertIsNone(database.claim_recipient(recipient_id, campaign_id))

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is installed with application dependencies")
    def test_iteration_report_builds_xlsx(self):
        campaign_id = database.create_campaign(["one@example.test"], "Subject", "Body", [])
        stream, filename = build_report(campaign_id, 1)
        self.assertTrue(filename.endswith("Iteration_1.xlsx"))
        self.assertGreater(len(stream.read()), 100)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is installed with application dependencies")
    def test_report_contains_recipient_audit_sheets_and_ist_timestamps(self):
        from openpyxl import load_workbook

        campaign_id = database.create_campaign(["sent@example.test", "failed@example.test"], "Subject", "Body", [])
        recipient_ids = database.queued_recipient_ids(campaign_id)
        for recipient_id, status, error in ((recipient_ids[0], "sent", None),
                                             (recipient_ids[1], "failed", "SMTP connection lost")):
            database.claim_recipient(recipient_id, campaign_id)
            database.record_attempt(recipient_id)
            database.recipient_result(recipient_id, status, error=error)

        stream, _ = build_report(campaign_id, 1)
        workbook = load_workbook(stream, read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["Campaign Summary", "Recipient Results", "Event Log"])
        summary = workbook["Campaign Summary"]
        self.assertEqual(summary.max_column, 15)
        recipients = workbook["Recipient Results"]
        self.assertEqual(recipients.max_row, 3)
        values = list(recipients.values)
        self.assertEqual([values[1][6], values[2][6]], ["SENT", "FAILED"])
        self.assertIn("PM IST", values[1][9])
        self.assertEqual(values[2][15], "SMTP_CONNECTION_LOST")
        self.assertGreater(workbook["Event Log"].max_row, 2)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
