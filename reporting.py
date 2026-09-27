"""Database-backed campaign audit workbooks; SMTP workers never write Excel."""

import re
from datetime import datetime, timezone, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import database

try:
    IST = ZoneInfo("Asia/Kolkata")
except Exception:  # Windows Python installations may not bundle the IANA database.
    IST = timezone(timedelta(hours=5, minutes=30), "IST")
TIME_FORMAT = "%d-%m-%Y %I:%M:%S %p IST"


def _timestamp(value):
    if not value:
        return "-"
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(IST).strftime(TIME_FORMAT)
    except (TypeError, ValueError, OverflowError):
        return "-"


def _duration(seconds):
    try:
        seconds = max(0, int(float(seconds or 0)))
    except (ValueError, TypeError, OverflowError):
        seconds = 0
    hours, remainder = divmod(seconds, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _failure_type(row):
    if row.get("error_type"):
        return row["error_type"]
    stage = (row.get("failure_stage") or "").upper()
    message = (row.get("error_message") or "").lower()
    if "timeout" in message:
        return "SMTP_TIMEOUT"
    if "disconnect" in message or "connection lost" in message or "closed" in message:
        return "SMTP_CONNECTION_LOST"
    if "auth" in message or stage == "LOGIN":
        return "SMTP_AUTH_ERROR"
    if "recipient" in message or stage == "RCPT_TO":
        return "SMTP_RECIPIENT_REJECTED"
    if stage in {"CONNECT", "EHLO", "STARTTLS", "MAIL_FROM", "RCPT_TO", "DATA"}:
        return "SMTP_SERVER_ERROR"
    return "UNKNOWN_ERROR" if row.get("error_message") else "-"


def _style(sheet, widths=None):
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="23415C")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for col in sheet.columns:
        letter = col[0].column_letter
        if widths and letter in widths:
            sheet.column_dimensions[letter].width = widths[letter]
        else:
            sheet.column_dimensions[letter].width = min(42, max(12, max(len(str(c.value or "")) for c in col) + 2))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {6, 17, 19})


def build_report(campaign_id, iteration_number=None, final=False):
    """Build a point-in-time report from the database without blocking sends."""
    from openpyxl import Workbook

    campaign, rows = database.iteration_results(campaign_id, iteration_number)
    if campaign is None:
        return None, None
    iteration = iteration_number or campaign["current_iteration"]
    events = campaign.pop("events", [])
    if final:
        rows = []
        events = []
        for number in range(1, campaign["current_iteration"] + 1):
            iteration_campaign, iteration_rows = database.iteration_results(campaign_id, number)
            rows.extend(iteration_rows)
            if iteration_campaign:
                events.extend(iteration_campaign.pop("events", []))
    counts = {key: sum(str(row["status"]).lower() == key for row in rows) for key in
              ("sent", "failed", "queued", "retrying", "sending", "cancelled")}
    start = campaign.get("started_at")
    completed = campaign.get("completed_at")
    duration_seconds = campaign.get("duration")
    if duration_seconds is None:
        try:
            end_dt = datetime.fromisoformat(completed) if completed else datetime.now(timezone.utc)
            start_dt = datetime.fromisoformat(start) if start else None
            duration_seconds = (end_dt - start_dt).total_seconds() if start_dt else 0
        except (TypeError, ValueError):
            duration_seconds = 0

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Campaign Summary"
    summary.append(["Campaign", "Campaign ID", "Iteration", "Status", "Start Time", "Scheduled Time",
                    "Last Updated", "Completion Time", "Total Recipients", "Sent", "Failed", "Pending",
                    "Retrying", "Paused", "Duration"])
    summary.append([campaign["name"], campaign_id, iteration, campaign["status"].upper(), _timestamp(start),
                    _timestamp(campaign.get("scheduled_at")), _timestamp(campaign.get("last_updated")),
                    _timestamp(completed), len(rows), counts["sent"], counts["failed"], counts["queued"],
                    counts["retrying"], counts["queued"] if campaign["status"] in ("paused", "auto_paused") else 0,
                    _duration(duration_seconds)])

    results = workbook.create_sheet("Recipient Results")
    results.append(["Campaign", "Campaign ID", "Iteration", "Email", "HR Name", "Company", "Status",
                    "First Attempt", "Last Attempt", "Sent At", "Failed At", "Attempt Count", "Retry Count",
                    "Sending Duration", "SMTP Status", "Error Type", "Error Message", "Last Event"])
    last_events = {}
    for event in events:
        if event.get("recipient_email"):
            last_events[(event["recipient_email"], event["iteration_number"])] = event["event"]
    # For final reports, preserve latest events across iterations as well.
    for row in rows:
        status = row["status"].upper()
        attempts = row.get("attempt_count") or 0
        error_type = _failure_type(row)
        last_event = last_events.get((row["recipient_email"], row["iteration_number"]),
                                     "SEND_SUCCESS" if status == "SENT" else "FINAL_FAILED" if status == "FAILED" else status)
        results.append([row["campaign_name"], campaign_id, row["iteration_number"], row["recipient_email"],
                        None, None, status, _timestamp(row.get("first_attempt_at")),
                        _timestamp(row.get("last_attempt_at")), _timestamp(row.get("sent_at")),
                        _timestamp(row.get("failed_at")), attempts, max(0, attempts - 1),
                        _duration(row.get("elapsed_seconds")),
                        row.get("smtp_response_code") or ("SUCCESS" if status == "SENT" else "-"),
                        error_type, row.get("error_message") or "-", last_event])

    event_sheet = workbook.create_sheet("Event Log")
    event_sheet.append(["Timestamp", "Campaign", "Campaign ID", "Iteration", "Email", "Event", "Status",
                        "Attempt", "Retry Count", "Details"])
    for event in events:
        event_sheet.append([_timestamp(event.get("occurred_at")), event.get("campaign_name"),
                            event.get("campaign_id"), event.get("iteration_number"),
                            event.get("recipient_email") or "-", event.get("event"),
                            (event.get("status") or "").upper(), event.get("attempt", 0),
                            event.get("retry_count", 0), event.get("details") or "-"])

    _style(summary, {"A": 30, "B": 38, "C": 12, "D": 18, "E": 28, "F": 28, "G": 28, "H": 28})
    _style(results, {"A": 28, "B": 38, "D": 32, "E": 22, "F": 24, "G": 14, "H": 28, "I": 28,
                     "J": 28, "K": 28, "N": 18, "O": 18, "P": 25, "Q": 48, "R": 24})
    _style(event_sheet, {"A": 28, "B": 28, "C": 38, "D": 12, "E": 32, "F": 28, "G": 16, "J": 48})
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", campaign["name"]).strip("_") or "Campaign"
    suffix = "Final" if final else f"Iteration_{iteration}"
    return stream, f"MailFlow_{safe_name}_{suffix}.xlsx"
